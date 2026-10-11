import re
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from sqlalchemy import func

from backend.utils.instance import db
from backend.communities.CommunityUserRole_models import CommunityUserRole
from backend.quests.subquest_completion import SubquestCompletion
from backend.models.models import Users
from backend.payments.wallet import ZecWallet
from backend.communities.community_models import CommunityUserXP, SprintUserXP
from backend.quests.sub_quest_models import Subquest
from backend.quests.quest_models import Quest  # ← adjust to wherever Quest lives

MAX_EXPORT_ROWS = 50_000


# ── ranking helpers ───────────────────────────────────────────────
def _banned_user_ids(community):
    return (
        db.session.query(CommunityUserRole.user_id)
        .filter(
            CommunityUserRole.community_id == community.id,
            CommunityUserRole.banned.is_(True),
        )
        .subquery()
    )


def ranked_sprint_users(sprint, community, n):
    """Top-n [(user_id, username, xp)] in the SAME order as api_sprint_leaderboard."""
    banned = _banned_user_ids(community)

    latest = (
        db.session.query(
            SubquestCompletion.user_id.label("user_id"),
            func.max(SubquestCompletion.completed_at).label("last_xp_at"),
        )
        .join(Subquest, SubquestCompletion.subquest_id == Subquest.id)
        .filter(
            Subquest.sprint_id == sprint.id,
            SubquestCompletion.status == "success",
            SubquestCompletion.completed_at.isnot(None),
        )
        .group_by(SubquestCompletion.user_id)
        .subquery()
    )

    rows = (
        db.session.query(Users.id, Users.username, SprintUserXP.xp)
        .select_from(SprintUserXP)
        .join(Users, Users.id == SprintUserXP.user_id)
        .outerjoin(latest, latest.c.user_id == Users.id)
        .filter(
            SprintUserXP.sprint_id == sprint.id,
            ~SprintUserXP.user_id.in_(banned),
        )
        .order_by(
            SprintUserXP.xp.desc(),
            latest.c.last_xp_at.asc().nullslast(),
            SprintUserXP.id.asc(),
        )
        .limit(n)
        .all()
    )
    return [(r.id, r.username, r.xp or 0) for r in rows]


def ranked_all_time_users(community, n):
    """Top-n [(user_id, username, xp)] in the SAME order as /api/<slug>/leaderboard."""
    banned = _banned_user_ids(community)

    latest = (
        db.session.query(
            SubquestCompletion.user_id.label("user_id"),
            func.max(SubquestCompletion.completed_at).label("last_xp_at"),
        )
        .join(Subquest, SubquestCompletion.subquest_id == Subquest.id)
        .join(Quest, Subquest.quest_id == Quest.id)
        .filter(
            Quest.community_id == community.id,
            SubquestCompletion.status == "success",
            SubquestCompletion.completed_at.isnot(None),
        )
        .group_by(SubquestCompletion.user_id)
        .subquery()
    )

    rows = (
        db.session.query(Users.id, Users.username, CommunityUserXP.xp)
        .select_from(CommunityUserXP)
        .join(Users, Users.id == CommunityUserXP.user_id)
        .outerjoin(latest, latest.c.user_id == Users.id)
        .filter(
            CommunityUserXP.community_id == community.id,
            ~CommunityUserXP.user_id.in_(banned),
        )
        .order_by(
            CommunityUserXP.xp.desc(),
            latest.c.last_xp_at.asc().nullslast(),
            CommunityUserXP.id.asc(),
        )
        .limit(n)
        .all()
    )
    return [(r.id, r.username, r.xp or 0) for r in rows]


# ── wallets ───────────────────────────────────────────────────────
def active_wallet_map(user_ids):
    """{user_id: shielded address}. If a user has several active wallets,
    the most recently connected one wins (ordered asc so later overwrites)."""
    if not user_ids:
        return {}

    wallets = (
        ZecWallet.query
        .filter(
            ZecWallet.user_id.in_(user_ids),
            ZecWallet.is_active.is_(True),
            # ZecWallet.verified.is_(True),   # uncomment to export verified wallets only
        )
        .order_by(ZecWallet.connected_at.asc())
        .all()
    )
    return {w.user_id: w.address for w in wallets}


# ── titles + filenames ────────────────────────────────────────────
def report_title(sprint):
    """'{Sprint Name} Ranking' for a sprint, 'Leaderboard Ranking' for all-time."""
    if sprint is not None:
        name = (sprint.title or "").strip()
        return f"{name} Ranking" if name else "Sprint Ranking"
    return "Leaderboard Ranking"


def export_filename(community, sprint):
    title = report_title(sprint)
    slug = re.sub(r"[^a-z0-9]+", "-", title.lower()).strip("-")
    if not slug or slug == "ranking":      # e.g. a sprint title with no latin characters
        slug = f"sprint-{str(sprint.uuid)[:8]}-ranking" if sprint else "leaderboard-ranking"
    return f"{community.slug}-{slug}.xlsx"


# ── xlsx ──────────────────────────────────────────────────────────
def _text_cell(ws, row, col, value):
    """Write as plain text. openpyxl turns strings starting with '=' into
    formulas, so a username like '=HYPERLINK(...)' would execute in Excel."""
    cell = ws.cell(row=row, column=col, value=value)
    cell.data_type = "s"
    return cell


def build_workbook(rows, wallets, title):
    wb = Workbook()
    ws = wb.active
    ws.title = "Ranking"

    # row 1: title (sprint titles are user input, so it goes in as plain text too)
    title_cell = _text_cell(ws, 1, 1, title)
    title_cell.font = Font(bold=True, size=14)
    title_cell.alignment = Alignment(horizontal="left", vertical="center")
    ws.merge_cells("A1:D1")
    ws.row_dimensions[1].height = 24

    # row 2 left blank, row 3: headers
    headers = ["No.", "Username", "XP", "Shielded address"]
    for col, label in enumerate(headers, start=1):
        c = ws.cell(row=3, column=col, value=label)
        c.font = Font(bold=True)
        c.alignment = Alignment(horizontal="left")

    for i, (user_id, username, xp) in enumerate(rows, start=1):
        r = i + 3
        ws.cell(row=r, column=1, value=i).alignment = Alignment(horizontal="left")
        _text_cell(ws, r, 2, username or "")
        # XP is a real number so it can be sorted / summed in Excel
        xp_cell = ws.cell(row=r, column=3, value=int(xp or 0))
        xp_cell.alignment = Alignment(horizontal="left")
        xp_cell.number_format = "#,##0"
        _text_cell(ws, r, 4, wallets.get(user_id, ""))

    ws.column_dimensions["A"].width = 8
    ws.column_dimensions["B"].width = 28
    ws.column_dimensions["C"].width = 12
    ws.column_dimensions["D"].width = 90
    ws.freeze_panes = "A4"

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf